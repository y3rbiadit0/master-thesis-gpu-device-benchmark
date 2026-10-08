#!/usr/bin/env python3
"""Build the defense from the supplied PowerPoint and retained thesis evidence.

Dependencies: python-pptx, Pillow. Run from any directory:
    python tools/build_defense.py

All plots are editable Office charts. The template is opened read-only, and its
masters, backgrounds, and attribution are retained. No measured data is rerun.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data/main-campaign"
OUT = ROOT / "presentation"
GREEN = "134E2D"
BRIGHT = "07902A"
LIME = "7EB730"
INK = "2E2E2E"
MUTED = "526258"
PALE = "EEF5EE"
WHITE = "FFFFFF"
TEAL = "277D8B"
ORANGE = "A65C20"
GRAY = "75827B"
SOURCES: dict[str, str] = {}
NOTES: list[dict] = []
CHARTS: list[dict] = []


def source(path: Path) -> str:
    SOURCES[str(path.relative_to(ROOT))] = hashlib.sha256(path.read_bytes()).hexdigest()
    return path.read_text()


def rows(name: str) -> list[dict]:
    return list(csv.DictReader(source(DATA / name).splitlines()))


def tex_rows(path: Path) -> list[list[str]]:
    return [[c.strip() for c in line.rstrip("\\").split("&")]
            for line in source(path).splitlines()
            if "&" in line and line.rstrip().endswith("\\\\")]


def rgb(c: str) -> RGBColor:
    return RGBColor.from_string(c)


def text(s, value, x, y, w, h, size=15, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font="Arial", valign=MSO_ANCHOR.TOP):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = f"Text: {str(value)[:65]}"
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = valign
    for i, line in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(3)
        p.line_spacing = 1.07
        p._p.get_or_add_pPr().append(OxmlElement("a:buNone"))
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
    return sh


def shape(s, x, y, w, h, fill=PALE, kind=MSO_SHAPE.ROUNDED_RECTANGLE,
          line=None):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.12
    # The imported theme otherwise gives even arrows and table cells a shadow.
    sh._element.spPr.append(OxmlElement("a:effectLst"))
    return sh


def arrow(s, x, y, w=.35, h=.18, color=BRIGHT):
    return shape(s, x, y, w, h, color, MSO_SHAPE.RIGHT_ARROW)


def chip(s, label, x, y, w=1.0, h=.62, fill=GREEN, size=16):
    shape(s, x, y, w, h, fill)
    text(s, label, x+.07, y+.03, w-.14, h-.06, size,
         WHITE if fill in (GREEN, BRIGHT, TEAL, INK, ORANGE, GRAY) else GREEN,
         True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def circle(s, label, x, y, d=.42, fill=GREEN, size=14):
    shape(s, x, y, d, d, fill, MSO_SHAPE.OVAL)
    text(s, label, x, y+.01, d, d-.02, size, WHITE, True,
         PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def card(s, heading, body, x, y, w, h, number=None, fill=PALE):
    shape(s, x, y, w, h, fill)
    offset = .15
    if number is not None:
        circle(s, str(number), x+.16, y+.18)
        offset = .72
    text(s, heading, x+offset, y+.16, w-offset-.15, .68, 17, GREEN, True)
    text(s, body, x+.17, y+.93, w-.34, h-1.03, 14, INK)


def head(s, title, section, subtitle=None):
    text(s, section.upper(), .65, .47, 7.8, .19, 9, MUTED, True)
    text(s, title, .65, .79, 8.7, .66, 30, GREEN, True, font="Cambria")
    if subtitle:
        text(s, subtitle, .66, 1.44, 8.68, .39, 13, MUTED)


def footer(s, index, citation="", appendix=False):
    # The template art occupies the outside corners; all content stays inset.
    text(s, citation, .65, 5.01, 8.02, .27, 8, GREEN)
    text(s, f"A{index-14}" if appendix else f"{index:02d} / 14",
         8.73, 5.01, .6, .22, 9, GREEN, align=PP_ALIGN.RIGHT)


def takeaway(s, value, y=4.48):
    sh = text(s, value, .65, y+.045, 8.70, .42, 13, GREEN)
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            r.font.italic = True


def note(s, title, seconds, script, sources, backup=False):
    n = len(NOTES)+1
    entry = {"slide": n, "title": title, "seconds": seconds,
             "backup": backup, "script": script, "sources": sources}
    NOTES.append(entry)
    prefix = "BACKUP — use during questions" if backup else f"TARGET: {seconds} seconds"
    s.notes_slide.notes_text_frame.text = (
        prefix + "\n\n" + script + "\n\nSOURCES / EVIDENCE\n" + "\n".join(sources)
    )


def legend(s, entries, x, y, step=2.12, size=10):
    for i, (label, color) in enumerate(entries):
        shape(s, x+i*step, y+.045, .12, .12, color, MSO_SHAPE.OVAL)
        text(s, label, x+i*step+.2, y, step-.23, .27, size, INK)


def chart(s, categories, series, x, y, w, h, title, kind="column",
          maximum=None, number_format="0.0", labels=False, log=False):
    """Native Office chart, including editable embedded workbook."""
    cd = CategoryChartData()
    cd.categories = categories
    for name, values, _ in series:
        cd.add_series(name, values)
    ct = XL_CHART_TYPE.LINE_MARKERS if kind == "line" else XL_CHART_TYPE.COLUMN_CLUSTERED
    ch = s.shapes.add_chart(ct, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.has_legend = False
    ch.has_title = True
    t = ch.chart_title.text_frame
    t.text = title
    for p in t.paragraphs:
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.bold = False
        p.font.color.rgb = rgb(MUTED)
    ch.font.name = "Arial"
    ch.font.size = Pt(11)
    ch.font.color.rgb = rgb(INK)
    for axis in [ch.category_axis, ch.value_axis]:
        axis.tick_labels.font.name = "Arial"
        axis.tick_labels.font.size = Pt(10)
        axis.tick_labels.font.color.rgb = rgb(INK)
        axis.major_tick_mark = XL_TICK_MARK.NONE
        axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.format.line.fill.background()
    ch.category_axis.has_major_gridlines = False
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = rgb("DDE5DD")
    ch.value_axis.major_gridlines.format.line.width = Pt(.5)
    ch.value_axis.minimum_scale = 0
    if maximum is not None:
        ch.value_axis.maximum_scale = maximum
    ch.value_axis.tick_labels.number_format = number_format
    ch.value_axis.tick_labels.number_format_is_linked = False
    if log:
        scale = ch.value_axis._element.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}scaling")
        b = OxmlElement("c:logBase")
        b.set("val", "10")
        scale.insert(0, b)
        ch.value_axis.minimum_scale = 1
    plot = ch.plots[0]
    if kind == "column":
        plot.gap_width = 80
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(10)
        dl.font.name = "Arial"
        dl.font.color.rgb = rgb(INK)
        dl.number_format = number_format
    for se, (_, _, color) in zip(ch.series, series):
        if kind == "line":
            se.format.line.color.rgb = rgb(color)
            se.format.line.width = Pt(2)
            se.marker.style = XL_MARKER_STYLE.CIRCLE
            se.marker.size = 5
            se.marker.format.fill.solid()
            se.marker.format.fill.fore_color.rgb = rgb(color)
            se.marker.format.line.fill.background()
            se.smooth = False
        else:
            se.format.fill.solid()
            se.format.fill.fore_color.rgb = rgb(color)
            se.format.line.fill.background()
    CHARTS.append({"slide": len(NOTES)+1, "title": title, "categories": categories,
                   "series": [{"name": n, "values": list(v)} for n, v, _ in series]})
    return ch


def simple_table(s, headers, body, widths, x=.65, y=1.8, row_h=.43, font=12):
    left = x
    for h, width in zip(headers, widths):
        shape(s, left, y, width, row_h, GREEN, MSO_SHAPE.RECTANGLE)
        text(s, h, left+.10, y+.09, width-.16, row_h-.12, font, WHITE, True)
        left += width
    for i, row in enumerate(body):
        left = x
        yy = y+(i+1)*row_h
        for item, width in zip(row, widths):
            shape(s, left, yy, width, row_h, WHITE if i % 2 else PALE, MSO_SHAPE.RECTANGLE)
            text(s, str(item), left+.10, yy+.085, width-.16, row_h-.10, font, INK)
            left += width


def prepare():
    prs = Presentation(ROOT / "thesis_style.pptx")
    # Final slide order is settled before any content is edited. All selected
    # slides already exist, so no unsafe XML-only slide duplication is needed.
    selection = [1, 12, 20, 14, 15, 13, 19, 17, 21, 22, 24, 23, 25, 28,
                 5, 8, 9, 18, 16, 26]
    ids = list(prs.slides._sldIdLst)
    chosen = [ids[i-1] for i in selection]
    for element in ids:
        prs.slides._sldIdLst.remove(element)
        if element not in chosen:
            prs.part.drop_rel(element.rId)
    for element in chosen:
        prs.slides._sldIdLst.append(element)

    def inset_art(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                inset_art(sh.shapes)
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.height < Inches(3):
                # These are the template's wave images (not the full-canvas
                # gradient). Group child coordinates use an identity transform.
                if sh.top < 0:
                    sh.top = int(sh.top*.30)
                    sh.height = int(sh.height*.30)
                elif sh.top > Inches(4):
                    sh.top = Inches(5.31)+int((sh.top-Inches(4.40))*.40)
                    sh.height = int(sh.height*.40)

    used_layouts = {s.slide_layout.part.partname: s.slide_layout for s in prs.slides}
    for layout in used_layouts.values():
        inset_art(layout.shapes)
        for sh in layout.shapes:
            if sh.has_text_frame and sh.text.startswith("CREDITS:"):
                sh.left, sh.top, sh.width, sh.height = map(Inches, [1.1, 4.96, 7.8, .35])
                sh.text_frame.margin_top = sh.text_frame.margin_bottom = 0
                for p in sh.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
                    for r in p.runs:
                        r.font.name = "Arial"
                        r.font.size = Pt(8)
                        r.font.color.rgb = rgb(GREEN)
    for i, slide in enumerate(prs.slides):
        for sh in list(slide.shapes):
            # Original vector waves on the cover are distinct from its layout.
            if i == 0 and sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                if sh.top > Inches(3):
                    sh.left, sh.top = Inches(7.1), Inches(5.03)
                    sh.width, sh.height = Inches(2.9), Inches(1.3)
                continue
            sh._element.getparent().remove(sh._element)
        for rel in list(slide.part.rels.values()):
            # Cover groups retain image-filled vector shapes; keep their image
            # relationships as well as the shapes themselves.
            retained_types = ("/slideLayout", "/notesSlide", "/image") if i == 0 else ("/slideLayout", "/notesSlide")
            if not rel.reltype.endswith(retained_types):
                slide.part.drop_rel(rel.rId)
        slide._element.attrib.pop("show", None)
    prs.core_properties.title = "GPU-Initiated Communication and Application Scalability"
    prs.core_properties.subject = "Master's thesis defense — 15 minutes"
    prs.core_properties.author = "Franco Nicolas Merenda"
    prs.core_properties.keywords = "GPU communication, NVSHMEM, NCCL, MPI, oneCCL, SYCL, CG"
    prs.core_properties.comments = "Built from thesis_style.pptx and retained thesis evidence."
    return prs


def build():
    OUT.mkdir(exist_ok=True)
    prs = prepare()
    halo = tex_rows(DATA / "halo_1d-bandwidth-table.tex")[1:]
    moe = {r[0]: [float(v) for v in r[1:]] for r in tex_rows(DATA / "moe-table.tex")[1:]}
    adapters = rows("adapter-comparison.csv")
    solver = rows("acg-summary.csv")
    sycl = tex_rows(DATA / "acg-sycl-table.tex")
    reduction = tex_rows(DATA / "reduction-error-table.tex")
    correlation = tex_rows(DATA / "cg-device-correlation-table.tex")
    preliminary = tex_rows(ROOT / "appendices/nccl-preliminary.tex")
    for path in [ROOT / "main.tex", ROOT / "chapters/01-introduction.tex",
                 ROOT / "chapters/04-methodology.tex", ROOT / "chapters/06-discussion.tex",
                 ROOT / "chapters/07-conclusion.tex", ROOT / "references.bib"]:
        source(path)

    def sv(matrix, backend, field="best_solver_s"):
        selected = sorted([r for r in solver if r["matrix"] == matrix and r["backend"] == backend],
                          key=lambda r: int(r["ranks"]))
        return [float(r[field]) for r in selected]

    # 01 — cover, using the template's original vector waves.
    s = prs.slides[0]
    text(s, "MASTER’S THESIS DEFENSE", 1.4, 1.14, 7.2, .25, 11, GREEN,
         True, PP_ALIGN.CENTER)
    text(s, "GPU-Initiated Communication\nand Application Scalability",
         1.10, 1.72, 7.8, 1.20, 33, GREEN, True, PP_ALIGN.CENTER, "Cambria")
    text(s, "From microbenchmarks to a conjugate gradient application",
         1.35, 3.06, 7.3, .54, 16, INK, align=PP_ALIGN.CENTER)
    text(s, "Franco Nicolas Merenda", 2.3, 3.77, 5.4, .35, 18,
         GREEN, True, PP_ALIGN.CENTER)
    text(s, "Università degli Studi di Salerno · Computer Science\n"
         "Supervisor: Prof. Biagio Cosenza · Co-supervisor: Lorenzo Carpentieri\n"
         "Academic year 2025–2026", 1.3, 4.30, 7.4, .75, 11, INK,
         align=PP_ALIGN.CENTER)
    note(s, "GPU-Initiated Communication and Application Scalability", 20,
         "My thesis investigates how communication affects applications distributed across GPUs. "
         "I compare communication paths, extend interfaces that expose them, and test whether the "
         "benchmark results survive in a complete conjugate gradient solver. The key question is "
         "when moving communication into GPU kernels actually helps the application.",
         ["main.tex: thesis metadata", "chapters/01-introduction.tex"])

    # 02 — motivation through two visual examples.
    s = prs.slides[1]
    head(s, "Fast GPUs still wait for data", "Motivation")
    text(s, "Scientific simulation", .83, 1.65, 3.9, .37, 19, GREEN, True)
    for i in range(3):
        chip(s, f"GPU {i}", .87+i*1.2, 2.23, .94, .66)
        for r in range(3):
            for c in range(3):
                shape(s, .92+i*1.2+c*.28, 3.05+r*.25, .24, .21,
                      LIME if c in [0, 2] else PALE, MSO_SHAPE.RECTANGLE)
        if i < 2:
            arrow(s, 1.84+i*1.2, 2.45, .23)
    text(s, "Exchange boundary values\nCombine small global sums", .86, 4.02, 3.75, .60, 15)
    text(s, "Distributed AI", 5.22, 1.65, 3.9, .37, 19, GREEN, True)
    chip(s, "Tokens", 5.30, 2.70, 1.12, .67, TEAL)
    for i, yy in enumerate([2.08, 2.85, 3.62]):
        line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(6.42), Inches(3.035),
                                      Inches(6.73), Inches(yy+.28))
        line.line.color.rgb = rgb(BRIGHT)
        line.line.width = Pt(1.4)
        chip(s, f"Expert {i+1}", 7.55, yy, 1.35, .55, GREEN, 14)
        arrow(s, 6.73, yy+.20, .79, .16)
    text(s, "Route activations across GPUs", 5.25, 4.26, 3.9, .32, 15)
    takeaway(s, "Scalability depends on time spent waiting, not just GPU compute speed.", 4.58)
    footer(s, 2, "Thesis Ch. 1; Hamidouche et al. (2025); Shen et al. (2026)")
    note(s, "Fast GPUs still wait for data", 50,
         "A GPU can execute local work very quickly, but a distributed application must also exchange data. "
         "On the left, a simulation divides its domain among GPUs. Each GPU needs values from neighboring "
         "partitions, and the algorithm also combines small scalar results across all participants. "
         "On the right, a mixture-of-experts model routes tokens to expert submodels held on different GPUs. "
         "These workloads have very different communication patterns, yet both can spend critical-path "
         "time waiting for data. That waiting is the motivation for looking beyond raw compute speed.",
         ["chapters/01-introduction.tex: Motivation",
          "Hamidouche et al., GPU-Initiated Networking for NCCL, arXiv:2511.15076",
          "Shen et al., Every Microsecond Matters, arXiv:2607.16100"])

    # 03 — distinguish initiation, data path, and request posting.
    s = prs.slides[2]
    head(s, "Move the call into the kernel", "Definition",
         "GPU-initiated communication: GPU threads issue the communication operation.")
    text(s, "Host-orchestrated path", .78, 2.01, 4, .30, 17, GREEN, True)
    for label, xx in [("Compute", .80), ("CPU submits", 3.05), ("Communicate", 5.30), ("Compute", 7.55)]:
        chip(s, label, xx, 2.48, 1.62, .57, GRAY if label == "CPU submits" else GREEN, 13)
        if xx < 7:
            arrow(s, xx+1.75, 2.68, .36)
    text(s, "Persistent GPU work", .78, 3.36, 4, .30, 17, GREEN, True)
    shape(s, .80, 3.87, 8.36, .70, PALE, line=BRIGHT)
    for i, label in enumerate(["Compute", "Issue transfer", "Wait / signal", "Consume"]):
        text(s, label, 1.0+i*2.08, 4.07, 1.5, .28, 14, GREEN, True, PP_ALIGN.CENTER)
        if i < 3:
            arrow(s, 2.58+i*2.08, 4.11, .30, .15)
    footer(s, 3, "GPU-aware memory ≠ GPU initiation. Device initiation can still use CPU-proxy network progress.")
    note(s, "Move the call into the kernel", 75,
         "GPU-initiated means that a GPU thread issues the communication operation from inside a kernel. "
         "In a conventional host-driven sequence, the CPU submits communication and arranges dependencies "
         "between producing and consuming kernels. Efficient host libraries can enqueue work asynchronously; "
         "this diagram is conceptual, not a claim that every host call blocks. A persistent implementation "
         "keeps repeated compute, communication, and synchronization inside device work. Two distinctions "
         "are essential. First, GPU-aware MPI can accept a device buffer without allowing a GPU thread "
         "to call MPI. Second, a GPU-issued request can still be handed to a CPU proxy for network posting. "
         "Direct GPU-memory data movement and direct GPU-to-NIC request posting are separate properties. "
         "The CPU also remains responsible for initialization and resource setup.",
         ["chapters/02-background.tex", "chapters/04-methodology.tex: NVSHMEM network posting"])

    # 04 — workload-driven popularity, without inventing adoption statistics.
    s = prs.slides[3]
    head(s, "Why the interest now?", "Motivation")
    card(s, "Smaller latency budgets", "Inference repeatedly waits on small collectives.", .65, 1.77, 2.70, 2.40, 1)
    card(s, "Data-dependent traffic", "MoE routing and irregular work need fine-grained control.", 3.65, 1.77, 2.70, 2.40, 2)
    card(s, "Better device interfaces", "Symmetric memory and device APIs enable fused execution.", 6.65, 1.77, 2.70, 2.40, 3)
    takeaway(s, "The opportunity: reduce orchestration and overlap useful work with communication.")
    footer(s, 4, "Hamidouche et al. (2025); Ma et al. (2026); Shen et al. (2026)")
    note(s, "Why the interest now?", 50,
         "The interest is driven by a combination of workloads and interfaces. During distributed "
         "inference, small collectives can lie directly on the token-generation critical path. "
         "Mixture-of-experts models also create routing-dependent exchanges that benefit from "
         "fine-grained control. At the same time, symmetric-memory runtimes and device-callable APIs "
         "make it practical to combine communication with computation. This creates an opportunity "
         "to reduce launches, host coordination, and exposed waiting. It does not guarantee a gain: "
         "the kernel must still move data, synchronize correctly, and share GPU resources with computation.",
         ["Hamidouche et al. (2025), arXiv:2511.15076",
          "Ma et al. (2026), Demystifying NVSHMEM, arXiv:2606.05951",
          "Shen et al. (2026), arXiv:2607.16100"])

    # 05 — ecosystem as capabilities rather than mutually exclusive brands.
    s = prs.slides[4]
    head(s, "One ecosystem, different contracts", "Background")
    simple_table(s, ["Path", "Who issues the call?", "What the programmer controls"], [
        ["GPU-aware MPI", "CPU", "Device buffers; message / collective semantics"],
        ["Traditional NCCL", "CPU, on a stream", "Collectives and grouped peer transfers"],
        ["NVSHMEM", "CPU or GPU", "Symmetric memory; puts, signals, ordering"],
        ["NCCL device API / GIN", "GPU kernel", "Registered windows; device-side operations"],
    ], [2.13, 2.17, 4.40], y=1.77, row_h=.46, font=12)
    shape(s, .65, 4.30, 8.7, .51, PALE)
    text(s, "OSHMPI: SHMEM over MPI     oneCCL: collective interface     SYCL: compute model",
         .8, 4.45, 8.4, .27, 12, GREEN)
    footer(s, 5, "NCCL device API introduced in 2.28; GIN from 2.28.7. NVIDIA documentation, checked 6 Oct 2026.")
    note(s, "One ecosystem, different contracts", 65,
         "MPI is the portable message-passing reference. GPU-aware MPI accepts device buffers, but the "
         "application calls it from the CPU. Traditional NCCL is also submitted from the host, usually "
         "on a CUDA stream, with optimized GPU communication kernels underneath. NVSHMEM exposes a "
         "symmetric-memory model with host and device operations: the programmer manages remote addressing, "
         "notification, and safe reuse. NCCL has recently added a user device API and GPU-Initiated "
         "Networking, or GIN. Its direct and proxy networking backends reinforce why invocation and "
         "progress must be distinguished. OSHMPI implements OpenSHMEM over MPI; it is not a GPU-initiated "
         "path in this study. oneCCL is the common communication interface that I extended, while SYCL "
         "is the compute programming model. These are different layers, not interchangeable categories.",
         ["chapters/02-background.tex",
          "NVIDIA, Device-Initiated Communication, current NCCL documentation, accessed 2026-10-06: "
          "https://docs.nvidia.com/deeplearning/nccl/user-guide/docs/usage/deviceapi.html",
          "Main campaign uses NCCL 2.18.5 and NVSHMEM 2.11.0; newer NCCL device APIs are context/future work."])

    # 06 — three research questions with an evidence ladder.
    s = prs.slides[5]
    head(s, "When does a faster path help?", "Research questions")
    for i, (h, b) in enumerate([
        ("Benefit", "When does device initiation repay its costs?"),
        ("Integration", "What do memory and completion adaptations cost?"),
        ("Transfer", "Which benchmark observations survive in a solver?"),
    ]):
        circle(s, str(i+1), .80, 1.87+i*.80, .45)
        text(s, h, 1.44, 1.85+i*.80, 1.65, .35, 18, GREEN, True)
        text(s, b, 3.06, 1.90+i*.80, 6.05, .41, 16)
    for i, label in enumerate(["Patterns", "Composite work", "Complete solver"]):
        chip(s, label, 1.03+i*3.02, 4.35, 2.34, .48, GREEN, 14)
        if i < 2:
            arrow(s, 3.60+i*3.02, 4.50, .34, .15)
    footer(s, 6, "Thesis: research questions and evaluation methodology")
    note(s, "When does a faster path help?", 35,
         "I organize the work around three questions. First, in which measured configurations does "
         "device initiation help? Second, what programming and runtime costs arise when these paths "
         "are exposed through a common interface? Third, which benchmark observations transfer to "
         "a complete application? The method is bottom-up: start with patterns, add composition, "
         "and finally include the solver's real dependencies and convergence.",
         ["chapters/01-introduction.tex: Research Questions",
          "chapters/04-methodology.tex: evaluation design"])

    # 07 — contributions and implementation status.
    s = prs.slides[6]
    head(s, "What I built and evaluated", "Contributions")
    card(s, "Unified benchmark", "CUDA + SYCL\nSeven configured stacks\nPrimitive + composite patterns",
         .65, 1.76, 2.70, 2.57)
    card(s, "oneCCL extensions", "Grouped NCCL point-to-point\nOSHMPI operation support\nNVSHMEM runtime foundation",
         3.65, 1.76, 2.70, 2.57)
    card(s, "Application extensions", "OSHMPI halo communicator\nSYCL implementation of aCG\nComplete-solver validation",
         6.65, 1.76, 2.70, 2.57)
    text(s, "Measured operation paths", .85, 4.49, 3.85, .29, 13, GREEN, True)
    text(s, "NVSHMEM adapter: runtime validated; operations unsupported", 4.32, 4.46, 4.78, .41, 11, MUTED)
    footer(s, 7, "Thesis Ch. 1: five contributions. Artifact references and implementation status in the notes.")
    note(s, "What I built and evaluated", 70,
         "The contributions connect measurement, interface engineering, and application validation. "
         "First, I built a common CUDA and SYCL benchmark harness covering seven configured stacks, "
         "with shared mapping, timing, and validation policies. Second, I extended oneCCL's NCCL "
         "backend so group boundaries reach NCCL, enabling mutually dependent batches of sends "
         "and receives. Third, I added the benchmarked OSHMPI operations and a separate NVSHMEM "
         "runtime foundation. That NVSHMEM adapter validates bootstrap, lifecycle, and symmetric "
         "staging, but its public operations remain unsupported and provide no performance results. "
         "Fourth, I added an OSHMPI halo communicator to the existing aCG suite. Finally, I implemented "
         "a SYCL version of aCG. The published native aCG solver and its device-NVSHMEM variant are "
         "the baseline application, not a new solver algorithm invented by this thesis.",
         ["chapters/01-introduction.tex: Contributions",
          "references.bib: merenda_gpu_comm_benchmark_2026; merenda_oneccl_nccl_groups_2026; "
          "merenda_oneccl_oshmpi_2026; merenda_oneccl_nvshmem_2026; merenda_acg_oshmpi_2026; merenda_acg_sycl_2026"])

    # 08 — topology and protocol.
    s = prs.slides[7]
    head(s, "Measure complete configured paths", "Experimental setup")
    text(s, "Leonardo Booster", .8, 1.74, 4.25, .38, 20, GREEN, True)
    text(s, "1–32 A100 GPUs · CUDA + SYCL\nNVLink within nodes · InfiniBand across nodes",
         .8, 2.23, 4.2, 1.0, 15)
    for j, xx in enumerate([5.20, 7.42]):
        shape(s, xx, 1.85, 1.9, 1.55, PALE, line=GREEN)
        text(s, f"Node {j+1}", xx+.16, 1.97, 1.5, .27, 13, GREEN, True)
        for k in range(4):
            chip(s, "GPU", xx+.16+(k%2)*.79, 2.43+(k//2)*.43, .65, .32, GREEN, 9)
    arrow(s, 7.12, 2.51, .25, .22)
    text(s, "Proxy-assisted NVSHMEM network posting", 5.18, 3.63, 4.12, .53, 14, GREEN, True)
    text(s, "Patterns → MoE / CG step → two sparse-matrix solves", .8, 3.40, 4.0, .86, 17, GREEN, True)
    takeaway(s, "8n4g = 8 nodes × 4 GPUs per node. GPU count alone does not describe the path.")
    footer(s, 8, "NCCL 2.18.5 · NVSHMEM 2.11.0 · common validation and completion-aware timing; details in A1.")
    note(s, "Measure complete configured paths", 55,
         "The main campaign runs on Leonardo's Booster partition, up to eight nodes with four A100 "
         "GPUs per node. NVLink and the inter-node network are different communication domains. "
         "I therefore label placement explicitly: eight-n-four-g means eight nodes and four GPUs "
         "per node. The measured NVSHMEM inter-node transport is ibrc, with CPU-proxy request posting; "
         "IBGDA was unavailable in this environment. Payloads can still use GPUDirect RDMA. The "
         "experiments progress from ping-pong, halos, reductions, and all-to-all through MoE and "
         "a composite CG step, then complete aCG solves on Bump and Queen. Comparisons describe "
         "configured paths, including their completion rules, rather than a pure change of call location.",
         ["chapters/04-methodology.tex: Experimental Setup, NVSHMEM network posting, benchmark protocol",
          "NVSHMEM collectives dispatch to NCCL in the default benchmark; aCG disables that dispatch."])

    # 09 — actual table values, not hand-restated chart data.
    s = prs.slides[8]
    head(s, "Persistence helps regular halos", "Finding 1 · GPU initiation")
    legend(s, [("CUDA MPI", TEAL), ("Device NVSHMEM", BRIGHT), ("OSHMPI", GRAY)], .82, 1.59, 2.18)
    cats = [re.search(r"\{(.*?)\}", r[0]).group(1) for r in halo]
    series = [("CUDA MPI", [float(r[1]) for r in halo], TEAL),
              ("Device NVSHMEM", [float(r[2]) for r in halo], BRIGHT),
              ("OSHMPI", [float(r[3]) for r in halo], GRAY)]
    chart(s, cats, series, .60, 1.98, 5.85, 2.30,
          "Steady halo · 16 MiB aggregate · effective GB/s ↑", maximum=50, labels=True)
    text(s, f"{float(halo[-1][2]):.1f}", 6.85, 2.10, 2.0, 1.06, 54, BRIGHT, True)
    text(s, "GB/s at 8n4g", 6.88, 3.10, 2.12, .35, 18, GREEN, True)
    text(s, "MPI: 35.7 GB/s\n100-exchange batches", 6.88, 3.66, 2.30, .64, 14)
    takeaway(s, "The favorable regime is sustained neighbor traffic; small / isolated exchanges favor MPI.")
    footer(s, 9, "Source: halo_1d-bandwidth-table.tex. Effective aggregate rate, not a single-link rate; different completion contracts.")
    note(s, "Persistence helps regular halos", 75,
         "The clearest favorable device-initiated regime is the steady halo. Here each sample contains "
         "one hundred regular neighbor exchanges, and NVSHMEM executes them inside a persistent GPU "
         "kernel with neighbor signaling. At eight nodes with four GPUs each, the effective aggregate "
         "rate is 40.6 gigabytes per second, compared with 35.7 for CUDA MPI and 24.4 for OSHMPI. "
         "The NVSHMEM rate is also nearly flat across the three multi-node placements. This is an "
         "aggregate benchmark rate, not the rate of one physical network link. MPI retains the "
         "small-message and isolated-exchange advantage. Importantly, the paths differ in batching, "
         "signaling, persistence, and completion as well as initiation. The result identifies a "
         "useful complete execution path; it does not isolate a GPU-initiation-only speedup.",
         ["data/main-campaign/halo_1d-bandwidth-table.tex",
          "chapters/04-methodology.tex: halo timing, 100-exchange batches",
          "chapters/06-discussion.tex: Device-Initiated Scaling"])

    # 10 — ratios calculated directly from generated MoE table.
    s = prs.slides[9]
    head(s, "Routing changes the winner", "Finding 2 · Workload sensitivity")
    legend(s, [("MPI", TEAL), ("NCCL", GREEN), ("NVSHMEM", BRIGHT), ("OSHMPI", GRAY)], .78, 1.58, 2.17)
    idx = [1, 3, 5]
    ms = [(name, [moe[name][i]/moe["MPI"][i] for i in idx], color)
          for name, color in [("MPI", TEAL), ("NCCL", GREEN), ("NVSHMEM", BRIGHT), ("OSHMPI", GRAY)]]
    chart(s, ["Uniform", "Locality-biased", "Hotspot 80%"], ms, .65, 1.96, 8.67, 2.37,
          "Dispatch + combine time / MPI · 8n4g · 32 MiB useful volume per rank · lower is better",
          maximum=2.1, number_format="0.00\"×\"", labels=True)
    takeaway(s, "Same total volume and topology; a different payload distribution changes the ranking.")
    footer(s, 10, "Source: moe-table.tex, selected direct paths. Routing, planning, and packing are outside the timed interval.")
    note(s, "Routing changes the winner", 65,
         "A uniform all-to-all does not describe every routed workload. This chart keeps topology and "
         "total useful volume fixed and changes only routing. Values are normalized to MPI within "
         "each routing case, so lower is better. Under uniform routing, NCCL takes about half MPI's "
         "time. With locality-biased routing, NCCL becomes 1.55 times slower than MPI, while NVSHMEM "
         "is favorable. With eighty percent of tokens targeting one rank, the two one-sided "
         "put-based paths take about 1.8 times MPI's time. This is a dispatch-and-combine "
         "microbenchmark; routing, planning, and packing are outside timing, and there is no "
         "end-to-end MoE model result. The lesson is to match the payload distribution, not "
         "just the collective name and byte count.",
         ["data/main-campaign/moe-table.tex (microseconds; 8n4g columns)",
          "All six measured stacks, uniform / locality / hotspot: " + "; ".join(
              f"{k}: {v[1]:.0f} / {v[3]:.0f} / {v[5]:.0f}" for k, v in moe.items()),
          "chapters/04-methodology.tex: MoE timing and useful-volume definition"])

    # 11 — adaptation changes memory and completion, not merely wrapper cost.
    s = prs.slides[10]
    head(s, "An adapter can change the path", "Finding 3 · Integration costs")
    a = next(r for r in adapters if r["pattern"] == "allreduce" and r["bytes_per_rank"] == "16777216")
    b = next(r for r in adapters if r["pattern"] == "alltoall")
    card(s, "oneCCL / NCCL", "Grouped peer transfers\nAsynchronous stream execution", .65, 1.80, 4.18, 1.55)
    card(s, "oneCCL / OSHMPI", "Symmetric staging buffers\nBlocking completion adaptation", 5.17, 1.80, 4.18, 1.55)
    text(s, f"{float(b['adapted_over_direct']):.2f}×", .85, 3.55, 1.80, .85, 42, GREEN, True)
    text(s, "adapted / direct time\nBulk all-to-all · 8n4g", 2.64, 3.70, 2.0, .64, 12)
    text(s, f"{float(a['adapted_over_direct']):.1f}×", 5.37, 3.55, 1.85, .85, 42, ORANGE, True)
    text(s, "adapted / direct time\n16 MiB allreduce · 1n4g", 7.28, 3.70, 1.97, .66, 12)
    takeaway(s, "The cost follows buffer movement and completion semantics, not the API name alone.")
    footer(s, 11, "Source: adapter-comparison.csv. Two different operation regimes; these are not initiation-only or wrapper-only comparisons.")
    note(s, "An adapter can change the path", 65,
         "The same common interface can be inexpensive or costly depending on the adaptations it "
         "requires. For the bulk all-to-all endpoint, oneCCL over NCCL takes 0.98 times direct "
         "NCCL's time: essentially comparable at this resolution. My grouped point-to-point "
         "extension preserves the batch structure and asynchronous execution. For a sixteen-mebibyte "
         "intra-node allreduce, oneCCL over OSHMPI takes 17.9 times direct OSHMPI's time, about "
         "6.54 milliseconds versus 0.364 milliseconds. Ordinary application buffers must be "
         "adapted to symmetric staging and a blocking completion contract. These are two "
         "different operation regimes, not a head-to-head ratio between adapters, and the "
         "comparisons also change compute interface or operand placement. They demonstrate "
         "the cost of the complete adapted path rather than pure wrapper overhead.",
         ["data/main-campaign/adapter-comparison.csv",
          "All-to-all endpoint: bytes_per_rank=8388608; direct 842.339 µs, adapted 824.939333 µs",
          "Allreduce endpoint: direct 364.465667 µs, adapted 6537.54 µs",
          "chapters/06-discussion.tex: RQ2"])

    # 12 — full application, both matrices, native charts from retained CSV.
    s = prs.slides[11]
    head(s, "The solver is the deciding test", "Finding 4 · Application transfer")
    legend(s, [("NCCL", GREEN), ("Device NVSHMEM", BRIGHT), ("OSHMPI halo + MPI sums", GRAY)], .80, 1.55, 2.73, 10)
    topologies = ["1n2g", "1n4g", "2n4g", "4n4g", "8n4g"]
    for j, matrix in enumerate(["Bump_2911", "Queen_4147"]):
        ss = [("NCCL", sv(matrix, "acg-cg-nccl"), GREEN),
              ("Device NVSHMEM", sv(matrix, "acg-device-nvshmem"), BRIGHT),
              ("OSHMPI", sv(matrix, "acg-cg-oshmpi"), GRAY)]
        chart(s, topologies, ss, .57+j*4.45, 1.91, 4.39, 2.19,
              f"{matrix} · solver time (s) ↓", kind="line", maximum=30 if j == 0 else 60,
              number_format="0")
        text(s, "8n4g: " + " / ".join(f"{values[-1]:.2f}" for _, values, _ in ss) + " s",
             .90+j*4.45, 4.16, 3.96, .26, 11, GREEN)
    takeaway(s, "Device NVSHMEM gains on blocking OSHMPI at scale; no universal GPU-initiated lead.")
    footer(s, 12, "Source: acg-summary.csv; fastest residual-valid trial per configuration. Full backend endpoints and medians: A1.")
    note(s, "The solver is the deciding test", 95,
         "The application test solves two large sparse systems with conjugate gradient. Each iteration "
         "combines local sparse work, halo exchange, scalar reductions, and a convergence dependency. "
         "These charts follow the thesis's representative fastest residual-valid trial per configuration; "
         "they are not median curves, and the appendix provides median endpoints as a check. "
         "Device-NVSHMEM improves relative to the blocking OSHMPI-halo solver at larger configurations. "
         "NCCL nevertheless has lower representative times across these retained curves. At eight "
         "nodes, for example, Queen takes 7.25 seconds with NCCL, 8.79 with device-NVSHMEM, and "
         "9.60 with OSHMPI. Bump's NCCL and device-NVSHMEM medians at that endpoint are very close, "
         "so I do not interpret that small gap as a robust ordering. The regular benchmark halo "
         "does not reproduce the solver's irregular peers and overlap. Matched scalar reductions "
         "agree more closely across levels. The meaningful target is exposed application time "
         "under the same result and completion requirements, not the fastest isolated transfer.",
         ["data/main-campaign/acg-summary.csv: best_solver_s; all plotted paths have 9 trials / 3 allocations per cell",
          "data/main-campaign/reduction-error-table.tex",
          "chapters/06-discussion.tex: Cross-Level Synthesis and Limitations",
          "At 8n4g, Bump medians: NCCL 6.369126 s; device NVSHMEM 6.413378 s."])

    # 13 — current research tied to concrete follow-up experiments.
    s = prs.slides[12]
    head(s, "What comes next?", "Recent advances → future work")
    for i, (left, right) in enumerate([
        ("NCCL device API + GIN", "Compare verified proxy and GPU-posted paths"),
        ("Symmetric collectives:\nwithin 7% of modeled bound", "Match initiation, memory, and completion"),
        ("Portable device communication", "Expose device handles through oneCCL / SYCL"),
    ]):
        yy = 1.88+i*.77
        chip(s, str(i+1), .78, yy, .43, .43, GREEN, 14)
        text(s, left, 1.40, yy+.02, 3.30, .64, 15, GREEN, True)
        arrow(s, 4.86, yy+.12, .35, .16)
        text(s, right, 5.46, yy+.01, 3.70, .57, 15)
    shape(s, .65, 4.35, 8.70, .49, PALE)
    text(s, "Then replay real peer graphs and validate the gain in complete applications.",
         .83, 4.49, 8.33, .28, 14, GREEN, True)
    footer(s, 13, "7%: Shen et al. (2026), modeled scale-up latency bound. NCCL: NVIDIA docs / Hamidouche et al. (2025).")
    note(s, "What comes next?", 80,
         "Recent interfaces make the next experiments especially relevant. NCCL now offers device "
         "communication and GIN, with both GPU-direct request-posting and proxy implementations. "
         "Recent work by Shen and colleagues builds low-latency collectives on NCCL's device API "
         "and reports overhead within seven percent of its modeled scale-up hardware lower bound. "
         "That is a published result in its studied regime, not a result of this thesis or a "
         "general scale-out guarantee. My first priority is a matched comparison of proxy and "
         "verified GPU-posted networking. Second, isolate initiation from persistence, allocation, "
         "batching, algorithm, and completion. Third, strengthen application correspondence by "
         "replaying real peer graphs and matching local work. Finally, newer NVSHMEM and NCCL "
         "device operations could be exposed through a portable oneCCL and SYCL interface. "
         "Every candidate improvement must then be tested in a complete application.",
         ["chapters/07-conclusion.tex: five future directions",
          "NVIDIA current NCCL documentation, accessed 2026-10-06: https://docs.nvidia.com/deeplearning/nccl/user-guide/docs/usage/deviceapi.html",
          "Hamidouche et al. (2025), arXiv:2511.15076, https://arxiv.org/abs/2511.15076",
          "Shen et al. (2026), arXiv:2607.16100, https://arxiv.org/abs/2607.16100; scale-up scope"])

    # 14 — keep the original Thanks slide and its inherited credit text.
    s = prs.slides[13]
    text(s, "Choose for the application", .83, .91, 8.34, .75,
         34, GREEN, True, PP_ALIGN.CENTER, "Cambria")
    for i, (a, b) in enumerate([
        ("Match the pattern", "Payloads, peers, topology, and local work"),
        ("Match completion", "Where the result is ready and who can use it"),
        ("Validate the solver", "Measure the waiting that remains exposed"),
    ]):
        circle(s, str(i+1), 1.13, 2.02+i*.62, .36, size=12)
        text(s, a, 1.68, 2.01+i*.62, 2.76, .36, 16, GREEN, True)
        text(s, b, 4.48, 2.04+i*.62, 4.25, .38, 13)
    text(s, "Thank you · Questions?", 2.0, 4.17, 6, .48, 23,
         GREEN, True, PP_ALIGN.CENTER, "Cambria")
    note(s, "Choose for the application", 40,
         "The conclusion is that GPU initiation is a conditional tool. It is effective when the "
         "traffic structure and execution model let it reduce exposed work, but the call location "
         "alone does not determine performance. My contribution is a connected method and set of "
         "software paths for testing that decision: match the application's communication pattern, "
         "match its completion requirements, and validate the complete solver. Thank you. "
         "I am happy to take questions.",
         ["chapters/07-conclusion.tex",
          "Presentation template: Greenwashing Impact Thesis Defense by Slidesgo. "
          "Original Thanks-slide attribution retained in the layout."])

    # A1 — complete solver endpoints and measurement qualifications.
    s = prs.slides[14]
    head(s, "Full solver endpoints & protocol", "Backup A1 · Evidence details",
         "8n4g · seconds · best / median residual-valid solver time")
    names = [("MPI", "acg-cg-mpi"), ("NCCL", "acg-cg-nccl"),
             ("Host NVSHMEM", "acg-cg-nvshmem"), ("OSHMPI halo + MPI", "acg-cg-oshmpi"),
             ("Device NVSHMEM", "acg-device-nvshmem")]
    body = []
    for label, back in names:
        body.append([label] + [f"{sv(mat, back)[-1]:.2f} / {sv(mat, back, 'median_solver_s')[-1]:.2f}"
                               for mat in ["Bump_2911", "Queen_4147"]])
    simple_table(s, ["Backend", "Bump_2911", "Queen_4147"], body,
                 [3.15, 2.76, 2.79], y=1.96, row_h=.38, font=12)
    text(s, "Solver: 3 allocations × 3 launches planned per configuration; completed valid runs retained.\n"
         "Patterns: median of allocation means; 5 allocations, or 3 for all-to-all and CG step.",
         .78, 4.44, 8.4, .50, 10, MUTED)
    footer(s, 15, "Sources: acg-summary.csv; thesis methodology. Best curves are representative minima, not typical-performance estimates.", True)
    note(s, "Full solver endpoints & protocol", 0,
         "This table exposes the other native backends and the difference between best and median "
         "statistics. Native MPI enters a slow mode in several multi-node configurations; the "
         "very large best-to-median spread must not be hidden by reporting only a fastest trial. "
         "At 8n4g, Bump MPI has eight retained trials; the other listed endpoint cells have nine, "
         "all across three allocations. Each solver run uses a manufactured solution and relative "
         "residual tolerance 1e-6. Timing covers solver execution, not initialization. Pattern "
         "centers are medians of allocation means, with trial means averaged within an allocation. "
         "Different backends have documented completion boundaries. The steady halo is a 100-exchange "
         "batch, not raw wire latency. NVSHMEM default benchmark collectives dispatch to NCCL; aCG "
         "disables that dispatch. The separate SYCL campaign uses NCCL 2.22.3 under oneCCL, unlike "
         "the main NCCL 2.18.5 campaign.",
         ["data/main-campaign/acg-summary.csv", "data/main-campaign/README.md",
          "chapters/04-methodology.tex: repetition, aggregation, timing, and aCG protocol"], True)

    # A2 — paired SYCL evidence, not cross-vendor performance portability.
    s = prs.slides[15]
    head(s, "SYCL can retain native performance", "Backup A2 · Portability")
    text(s, "0.87–1.02×", .90, 1.86, 5.0, 1.0, 48, GREEN, True)
    text(s, "Paired SYCL / native solver-time ratios\nWithin one node · both matrices", .96, 3.00, 4.7, .70, 17)
    card(s, "Keep the execution contract", "Match layout and synchronization.\nAcross nodes, host completion can dominate.",
         5.88, 1.93, 3.40, 2.41)
    takeaway(s, "This is performance retention on NVIDIA A100, not cross-vendor performance portability.")
    footer(s, 16, "Source: acg-sycl-table.tex; paired-ratio aggregation. Separate allocations and versions from the main campaign.", True)
    note(s, "SYCL can retain native performance", 0,
         "The SYCL study interleaves native and SYCL executions in the same allocations. The "
         "reported ratio is the median over allocations of per-allocation ratios of medians, "
         "not necessarily the quotient of the two global medians. Across one to four GPUs, "
         "the reported ratios are 0.87 to 1.02 for the two matrices. Matching 32-bit sparse "
         "indices, device-resident scalars, and one host convergence wait per iteration is "
         "important. Across nodes the oneCCL/NCCL path imposes host completion on reductions, "
         "and it must not be interpreted as a programming-model-only cost. The paired NCCL "
         "comparison also changes library version. The MPI slow mode in the native path is "
         "not a universal property of MPI: the SYCL MPI path continues scaling on the same "
         "installation. All claims concern the measured NVIDIA platform.",
         ["data/main-campaign/acg-sycl-table.tex",
          "data/main-campaign/acg-sycl-comparison.csv",
          "chapters/06-discussion.tex: SYCL Portability Cost"], True)

    # A3 — what benchmark correspondence does and does not establish.
    s = prs.slides[16]
    head(s, "Transfer requires matched dependencies", "Backup A3 · Benchmark correspondence")
    card(s, "Matched scalar reductions", "MPI / NCCL, one FP64 sum\n\nMean relative differences:\n7.8–11.1% across five topologies", .65, 1.85, 4.18, 2.30)
    card(s, "Device CG-step scaling", "Grid side 8192: 8.0% / 12.6%\nGrid side 512: 582% / 853%\n\nBump / Queen mean deviations", 5.17, 1.85, 4.18, 2.30)
    takeaway(s, "Even an unchanged communication sequence can mislead when local work per GPU differs.")
    footer(s, 17, "Sources: reduction-error-table.tex; cg-device-correlation-table.tex. Descriptive checks, not held-out predictions.", True)
    note(s, "Transfer requires matched dependencies", 0,
         "Two narrowly matched checks support correspondence. The matched MPI and NCCL scalar "
         "intervals differ by 7.8 to 11.1 percent on average across five configurations. The "
         "device CG-step scaling shape agrees with the monolithic solver at grid side 8192 "
         "but differs enormously at side 512, even though the communication sequence and "
         "backend are unchanged. Local work per GPU and the work-to-transfer ratio are therefore "
         "essential matching variables. The regular two-neighbor benchmark halo does not "
         "capture the solver's directed peer graph or its overlap. These are descriptive "
         "comparisons, not trained models or held-out predictions. The scalar comparison "
         "uses MPI reductions from the OSHMPI-halo solver and reductions from the NCCL solver.",
         ["data/main-campaign/reduction-error-table.tex: " + str(reduction[1:]),
          "data/main-campaign/cg-device-correlation-table.tex: " + str(correlation[1:]),
          "chapters/06-discussion.tex: Cross-Level Synthesis"], True)

    # A4 — preliminary NCCL study, explicit scope and regressions.
    s = prs.slides[17]
    head(s, "Newer NCCL: a follow-up candidate", "Backup A4 · Exploratory evidence",
         "1n4g · supplied three-run medians · allreduce latency in µs ↓")
    allr = preliminary[1:6]
    ps = [("2.18.5 ordinary", [float(r[2]) for r in allr[:3]], GRAY),
          ("2.31.2 ordinary", [float(r[3]) for r in allr[:3]], TEAL),
          ("2.31.2 symmetric", [float(r[4]) for r in allr[:3]], BRIGHT)]
    legend(s, [(p[0], p[2]) for p in ps], .82, 1.96, 2.75)
    chart(s, ["8 B", "32 KiB", "256 KiB"], ps, .65, 2.28, 5.47, 2.02,
          "Allreduce latency (µs)", maximum=28, labels=True)
    text(s, "Gain at small sizes\nRegression at 256 KiB", 6.44, 2.65, 2.83, .83, 17, GREEN, True)
    text(s, "Small all-to-all does not\nshare the small-size gain.", 6.44, 3.70, 2.83, .52, 13)
    takeaway(s, "Symmetric allocation / registration does not establish device-side invocation.")
    footer(s, 18, "Source: appendix, nccl-preliminary.tex. Invocation, datatype, full timing contract, and individual trials unavailable.", True)
    note(s, "Newer NCCL: a follow-up candidate", 0,
         "This exploratory study is outside the main campaign and does not answer the research "
         "questions. The eight-byte allreduce median changes from 12.4 to 7.5 microseconds with "
         "symmetric allocation and registration in NCCL 2.31.2, but ordinary allocation in the "
         "new version remains at 12.6. At 256 KiB the symmetric result is slower. Small all-to-all "
         "also does not show the allreduce benefit. The eight-byte payload is not automatically "
         "one FP64 sum because datatype is unrecorded. Exact source, launch configuration, "
         "completion, warm-up, validation output, and individual trials are unavailable. No "
         "uncertainty interval or application speedup can be reconstructed. The next step is "
         "a controlled reduction with verified invocation and GPU-consumable completion.",
         ["appendices/nccl-preliminary.tex", "NVIDIA NCCL buffer registration and device API documentation"], True)

    # A5 — threats linked to executable experiments.
    s = prs.slides[18]
    head(s, "What the evidence leaves open", "Backup A5 · Limits & next experiments")
    simple_table(s, ["Current limit", "Experiment that addresses it"], [
        ["One cluster; proxy-assisted networking", "Matched proxy / verified GPU-posted runs"],
        ["Initiation changes with execution structure", "Vary invocation, persistence, and batching separately"],
        ["Regular halos; two solver matrices", "Replay real peer graphs; expand application set"],
        ["Older measured library paths", "Compare newer NVSHMEM and NCCL device APIs"],
        ["Portable adapter completion / staging", "Device handles + stream-ordered completion"],
    ], [4.07, 4.63], y=1.83, row_h=.40, font=11)
    takeaway(s, "Report complete-path gains and losses; do not infer a universal library ordering.")
    footer(s, 19, "Thesis discussion and conclusions. Few independent allocations; no hypothesis tests; MoE is a microbenchmark only.", True)
    note(s, "What the evidence leaves open", 0,
         "The archive characterizes one platform and configured paths. GPU initiation was not "
         "varied independently of persistence, batching, signaling, and kernel organization. "
         "There are two sparse matrices and no end-to-end MoE application. Sampling uses a "
         "small number of independent allocations, without hypothesis tests, so small point "
         "differences are not robust rankings. Future experiments should isolate those factors, "
         "verify transport selection and GPU/NIC placement, measure equivalent GPU-visible "
         "result states, and test complete applications. Portable device communication would "
         "need explicit ownership for symmetric windows and device handles without losing "
         "the host interface's stream and completion semantics.",
         ["chapters/06-discussion.tex: Limitations", "chapters/07-conclusion.tex: Future Work"], True)

    # A6 — readable references, artifact pointers in notes.
    s = prs.slides[19]
    head(s, "Selected sources & artifacts", "Backup A6 · References")
    refs = [
        ("NVSHMEM mechanisms", "Ma et al. (2026). Demystifying NVSHMEM. arXiv:2606.05951."),
        ("NCCL networking", "Hamidouche et al. (2025). GPU-Initiated Networking for NCCL. arXiv:2511.15076."),
        ("Latency-bound collectives", "Shen et al. (2026). Every Microsecond Matters. arXiv:2607.16100."),
        ("API definitions", "NVIDIA. NCCL Device-Initiated Communication documentation. Checked 6 Oct 2026."),
        ("Measured evidence", "Retained main-campaign tables, solver summaries, and thesis methodology."),
        ("Contributed software", "Benchmark suite; three oneCCL extensions; OSHMPI and SYCL aCG extensions."),
    ]
    for i, (h, b) in enumerate(refs):
        yy = 1.77+i*.49
        circle(s, str(i+1), .76, yy+.03, .26, GREEN, 9)
        text(s, h, 1.19, yy, 2.25, .29, 12, GREEN, True)
        text(s, b, 3.52, yy, 5.70, .42, 10)
    footer(s, 20, "Full citations, artifact identifiers, and data provenance are in the speaker notes and presentation/source_manifest.json.", True)
    bib = source(ROOT / "references.bib")
    artifact_keys = ["merenda_gpu_comm_benchmark_2026", "merenda_oneccl_nccl_groups_2026",
                     "merenda_oneccl_oshmpi_2026", "merenda_oneccl_nvshmem_2026",
                     "merenda_acg_oshmpi_2026", "merenda_acg_sycl_2026"]
    artifact_refs = []
    for key in artifact_keys:
        m = re.search(r"@\w+\{" + re.escape(key) + r",(.*?)(?=\n@|\Z)", bib, re.S)
        if m:
            url = re.search(r"url\s*=\s*\{([^}]+)\}", m.group(1))
            artifact_refs.append(key + (": " + url.group(1) if url else ""))
    note(s, "Selected sources & artifacts", 0,
         "These references support the background and recent-developments slides. External results "
         "are attributed to their authors; the thesis measurements use the retained archive "
         "identified in the accompanying manifest. The three recent papers are preprints in "
         "the cited arXiv versions. The native aCG baseline is credited in the thesis to "
         "Trotter et al. (2025); this thesis extends its communication and compute paths. "
         "The presentation is derived from the supplied Slidesgo template, and its original "
         "Thanks-slide credits are retained.",
         ["https://arxiv.org/abs/2606.05951", "https://arxiv.org/abs/2511.15076",
          "https://arxiv.org/abs/2607.16100",
          "https://docs.nvidia.com/deeplearning/nccl/user-guide/docs/usage/deviceapi.html",
          "references.bib: trotter_cpu-_2025"] + artifact_refs, True)

    assert len(prs.slides) == len(NOTES) == 20
    assert sum(n["seconds"] for n in NOTES) == 840
    # Native chart bounds and actual textboxes must remain inside the canvas.
    for n, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:  # original cover bleed
                continue
            assert sh.left >= 0 and sh.top >= 0, (n, sh.name)
            assert sh.left+sh.width <= prs.slide_width+10, (n, sh.name)
            assert sh.top+sh.height <= prs.slide_height+10, (n, sh.name)
    path = ROOT / "thesis_defense.pptx"
    prs.save(path)
    manuscript = ["# Thesis defense — speaker notes", "",
                  "14 main slides · target 14:00 · one-minute buffer before questions.", "",
                  "Slides 15–20 are technical backup slides. Stop the main talk on slide 14.", ""]
    elapsed = 0
    for n in NOTES:
        time = "Backup" if n["backup"] else f"{n['seconds']} s · start {elapsed//60:02d}:{elapsed%60:02d}"
        manuscript += [f"## {n['slide']}. {n['title']}", "", f"*{time}*", "", n["script"], "", "Sources:"]
        manuscript += [f"- {v}" for v in n["sources"]] + [""]
        elapsed += n["seconds"]
    (OUT / "speaker_notes.md").write_text("\n".join(manuscript)+"\n")
    manifest = {"template": "thesis_style.pptx", "template_sha256": hashlib.sha256((ROOT / "thesis_style.pptx").read_bytes()).hexdigest(),
                "main_slides": 14, "backup_slides": 6, "target_seconds": 840,
                "evidence_sha256": SOURCES, "charts": CHARTS,
                "external_context_checked": "2026-10-06",
                "note": "Charts use retained evidence. External context is not a new campaign measurement."}
    (OUT / "source_manifest.json").write_text(json.dumps(manifest, indent=2)+"\n")
    print(f"Created {path}: 14 main + 6 backup slides; {len(CHARTS)} native charts; 14:00 speaking target.")


if __name__ == "__main__":
    build()
